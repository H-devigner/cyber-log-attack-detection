"""Build the project presentation deck."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_PATH = PROJECT_ROOT / "presentation" / "cyber_log_attack_detection_overview.pptx"

COLORS = {
    "ink": RGBColor(28, 35, 43),
    "muted": RGBColor(88, 99, 112),
    "blue": RGBColor(42, 111, 151),
    "red": RGBColor(204, 80, 72),
    "green": RGBColor(61, 142, 109),
    "amber": RGBColor(231, 166, 70),
    "light": RGBColor(247, 249, 250),
    "line": RGBColor(210, 216, 222),
}


def add_title(slide, title: str, subtitle: str | None = None) -> None:
    box = slide.shapes.add_textbox(Inches(0.55), Inches(0.35), Inches(12.2), Inches(0.6))
    paragraph = box.text_frame.paragraphs[0]
    paragraph.text = title
    paragraph.font.size = Pt(28)
    paragraph.font.bold = True
    paragraph.font.color.rgb = COLORS["ink"]
    if subtitle:
        sub = slide.shapes.add_textbox(Inches(0.58), Inches(0.95), Inches(11.7), Inches(0.35))
        p = sub.text_frame.paragraphs[0]
        p.text = subtitle
        p.font.size = Pt(12)
        p.font.color.rgb = COLORS["muted"]


def add_footer(slide, slide_no: int) -> None:
    line = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(7.17), Inches(13.33), Inches(0.04))
    line.fill.solid()
    line.fill.fore_color.rgb = COLORS["line"]
    line.line.fill.background()
    box = slide.shapes.add_textbox(Inches(0.55), Inches(7.22), Inches(12.2), Inches(0.25))
    p = box.text_frame.paragraphs[0]
    p.text = f"Cyber Log Attack Detection | {slide_no}"
    p.font.size = Pt(8)
    p.font.color.rgb = COLORS["muted"]


def add_bullets(slide, items: list[str], x: float, y: float, w: float, h: float, size: int = 18) -> None:
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    for idx, item in enumerate(items):
        paragraph = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        paragraph.text = item
        paragraph.level = 0
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = COLORS["ink"]
        paragraph.space_after = Pt(8)


def add_card(slide, title: str, body: str, x: float, y: float, w: float, h: float, color: RGBColor) -> None:
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = COLORS["light"]
    shape.line.color.rgb = color
    shape.line.width = Pt(1.5)

    title_box = slide.shapes.add_textbox(Inches(x + 0.18), Inches(y + 0.12), Inches(w - 0.36), Inches(0.28))
    p = title_box.text_frame.paragraphs[0]
    p.text = title
    p.font.bold = True
    p.font.size = Pt(14)
    p.font.color.rgb = color

    body_box = slide.shapes.add_textbox(Inches(x + 0.18), Inches(y + 0.48), Inches(w - 0.36), Inches(h - 0.55))
    bp = body_box.text_frame.paragraphs[0]
    bp.text = body
    bp.font.size = Pt(11)
    bp.font.color.rgb = COLORS["ink"]


def add_table(slide, headers: list[str], rows: list[list[str]], x: float, y: float, w: float, h: float) -> None:
    table_shape = slide.shapes.add_table(len(rows) + 1, len(headers), Inches(x), Inches(y), Inches(w), Inches(h))
    table = table_shape.table
    for col_idx, header in enumerate(headers):
        cell = table.cell(0, col_idx)
        cell.text = header
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLORS["blue"]
        for paragraph in cell.text_frame.paragraphs:
            paragraph.font.color.rgb = RGBColor(255, 255, 255)
            paragraph.font.bold = True
            paragraph.font.size = Pt(10)
    for row_idx, row in enumerate(rows, start=1):
        for col_idx, value in enumerate(row):
            cell = table.cell(row_idx, col_idx)
            cell.text = value
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(255, 255, 255) if row_idx % 2 else COLORS["light"]
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.size = Pt(9)
                paragraph.font.color.rgb = COLORS["ink"]


def add_metric_chart(slide, x: float, y: float, w: float, h: float) -> None:
    chart_data = CategoryChartData()
    chart_data.categories = ["Binary", "Category", "Binary baseline", "Category baseline"]
    chart_data.add_series("Macro F1", (0.9929, 0.9447, 0.3371, 0.0507))
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
        chart_data,
    ).chart
    chart.has_legend = False
    chart.value_axis.maximum_scale = 1.0
    chart.value_axis.minimum_scale = 0.0
    chart.value_axis.tick_labels.font.size = Pt(9)
    chart.category_axis.tick_labels.font.size = Pt(9)
    plot = chart.plots[0]
    plot.has_data_labels = True
    plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    plot.data_labels.number_format = "0.00"
    plot.data_labels.font.size = Pt(8)
    chart.series[0].format.fill.solid()
    chart.series[0].format.fill.fore_color.rgb = COLORS["green"]


def add_source_chart(slide, x: float, y: float, w: float, h: float) -> None:
    chart_data = CategoryChartData()
    chart_data.categories = ["Firewall", "Web", "SSH"]
    chart_data.add_series("Rows", (27000, 24000, 20000))
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.BAR_CLUSTERED,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
        chart_data,
    ).chart
    chart.has_legend = False
    chart.value_axis.tick_labels.font.size = Pt(9)
    chart.category_axis.tick_labels.font.size = Pt(10)
    plot = chart.plots[0]
    plot.has_data_labels = True
    plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    plot.data_labels.number_format = "#,##0"
    plot.data_labels.font.size = Pt(8)
    chart.series[0].format.fill.solid()
    chart.series[0].format.fill.fore_color.rgb = COLORS["blue"]


def blank_slide(prs: Presentation, title: str, subtitle: str | None = None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, title, subtitle)
    add_footer(slide, len(prs.slides))
    return slide


def build_deck() -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide = blank_slide(prs, "Cyber Log Attack Detection", "SSH, web-server, and firewall attack detection with machine learning")
    hero = slide.shapes.add_textbox(Inches(0.85), Inches(2.0), Inches(11.8), Inches(1.1))
    p = hero.text_frame.paragraphs[0]
    p.text = "A two-pass ML pipeline: first detect attack vs normal, then classify attack type with source-specific specialists."
    p.font.size = Pt(26)
    p.font.bold = True
    p.font.color.rgb = COLORS["ink"]
    add_card(slide, "Data", "Public logs + synthetic coverage for missing labels and scenarios", 0.85, 3.65, 3.7, 1.25, COLORS["blue"])
    add_card(slide, "Model", "Unified binary detector plus SSH/web/firewall category specialists", 4.85, 3.65, 3.7, 1.25, COLORS["green"])
    add_card(slide, "Demo", "Scenario batch plus Podman ELK ingestion with Kibana", 8.85, 3.65, 3.7, 1.25, COLORS["amber"])

    slide = blank_slide(prs, "Problem And Goal")
    add_bullets(
        slide,
        [
            "Security teams receive noisy SSH, web, and firewall logs.",
            "A single log line is not always enough; context and source type matter.",
            "Goal: build a repeatable data-science pipeline from logs to EDA, modeling, evaluation, and demo scoring.",
            "Output: normal vs attack first, attack category second.",
        ],
        0.9,
        1.45,
        11.7,
        4.7,
    )

    slide = blank_slide(prs, "Data Sources")
    add_table(
        slide,
        ["Type", "Public source", "Synthetic?", "Label caveat"],
        [
            ["Firewall", "UCI firewall data", "Yes", "Action field used as weak label"],
            ["Web", "NASA HTTP access logs", "Yes", "Public rows normal; attacks simulated"],
            ["SSH", "Zenodo SSH honeypot", "Yes", "Attack-heavy; normal SSH simulated"],
        ],
        0.7,
        1.35,
        12.0,
        2.0,
    )
    add_source_chart(slide, 1.15, 3.65, 10.6, 2.65)

    slide = blank_slide(prs, "Why Synthetic Data Was Needed")
    add_card(slide, "Coverage gap", "Public logs rarely cover normal and attack behavior with clean labels in one place.", 0.75, 1.35, 3.8, 1.35, COLORS["red"])
    add_card(slide, "Class coverage", "Synthetic rows ensure every demo category has enough examples to train and evaluate.", 4.75, 1.35, 3.8, 1.35, COLORS["amber"])
    add_card(slide, "Prototype speed", "It lets the pipeline be built before real organization logs are available.", 8.75, 1.35, 3.8, 1.35, COLORS["green"])
    add_bullets(
        slide,
        [
            "Synthetic data is useful for bootstrapping, not final proof.",
            "Production readiness needs real logs and analyst-reviewed labels.",
        ],
        1.0,
        3.55,
        11.2,
        1.4,
        size=20,
    )

    slide = blank_slide(prs, "Normalized Feature Table")
    add_bullets(
        slide,
        [
            "All sources are converted into one 46-column table.",
            "Common features: log source, event type, protocol, service, timestamp-derived fields.",
            "SSH features: auth result, username type, failed login count, success after failures.",
            "Web features: method, status family, URL length, query length, suspicious tokens, user-agent family.",
            "Firewall features: action flags, ports, service, packet/byte volume, source event count.",
        ],
        0.85,
        1.25,
        12.0,
        5.3,
        size=17,
    )

    slide = blank_slide(prs, "Window Aggregation")
    add_table(
        slide,
        ["Source", "Current state", "Production recommendation"],
        [
            ["SSH", "Partial: ssh_failed_logins_10m", "5-10 min per source IP/user"],
            ["Web", "Mostly per request", "5-10 min for scanners; one request can flag SQLi/XSS"],
            ["Firewall", "Partial: src_event_count", "5 min port scan; 10-30 min outbound context"],
        ],
        0.8,
        1.35,
        11.7,
        2.2,
    )
    add_bullets(
        slide,
        [
            "Current model is tabular, not a sequence model.",
            "Next major improvement: compute real rolling windows before scoring.",
        ],
        1.0,
        4.15,
        11.0,
        1.1,
        size=20,
    )

    slide = blank_slide(prs, "Model Architecture")
    add_card(slide, "Pass 1", "Unified binary detector predicts normal vs attack across all log sources.", 0.9, 1.45, 3.7, 1.3, COLORS["blue"])
    add_card(slide, "Routing", "Only attack rows continue. log_source chooses ssh, web, or firewall specialist.", 4.85, 1.45, 3.7, 1.3, COLORS["amber"])
    add_card(slide, "Pass 2", "Specialist category detector predicts the attack type.", 8.8, 1.45, 3.7, 1.3, COLORS["green"])
    add_bullets(
        slide,
        [
            "Model: SGDClassifier with logistic loss, elastic-net regularization, balanced class weights, and early stopping.",
            "Split: 80/20 stratified train/test; 10% internal validation for early stopping.",
            "Baselines: most-frequent-class dummy models.",
        ],
        0.95,
        3.6,
        11.7,
        2.1,
        size=17,
    )

    slide = blank_slide(prs, "Results")
    add_metric_chart(slide, 0.85, 1.25, 6.25, 4.5)
    add_table(
        slide,
        ["Decision", "Evidence"],
        [
            ["Binary", "Unified binary macro F1: 0.9929"],
            ["Category", "Unified category macro F1: 0.9447"],
            ["Specialists", "Web and firewall category specialists improved diagnosis"],
            ["Final", "Unified binary + specialist categories"],
        ],
        7.35,
        1.45,
        5.3,
        2.75,
    )
    add_bullets(slide, ["Metrics prove pipeline behavior on current data, not production readiness."], 7.45, 4.55, 5.0, 0.8, size=16)

    slide = blank_slide(prs, "Scenario Demo")
    add_table(
        slide,
        ["Metric", "Value"],
        [
            ["Rows scored", "18"],
            ["Expected attacks", "12"],
            ["Predicted attacks", "11"],
            ["Binary accuracy", "0.944"],
            ["Category accuracy on attacks", "0.833"],
        ],
        0.9,
        1.35,
        5.2,
        2.55,
    )
    add_bullets(
        slide,
        [
            "One SQL injection was categorized as XSS.",
            "One suspicious outbound firewall event was missed by the binary gate.",
            "These visible mistakes are useful: they show where more real data and better windows are needed.",
        ],
        6.7,
        1.45,
        5.8,
        2.65,
        size=17,
    )

    slide = blank_slide(prs, "ELK With Podman")
    add_bullets(
        slide,
        [
            "Logstash parses sample SSH, web, and firewall logs.",
            "Elasticsearch stores normalized events.",
            "Kibana supports analyst inspection and dashboards.",
            "Podman starts the local stack with: podman-compose -f compose.yaml up -d",
            "ELK is ingestion/exploration; ML scoring is a separate pipeline.",
        ],
        0.9,
        1.35,
        11.8,
        4.4,
        size=18,
    )

    slide = blank_slide(prs, "Limitations")
    add_table(
        slide,
        ["Risk", "Why it matters", "Fix"],
        [
            ["Synthetic data", "Can be too clean", "Add real labeled logs"],
            ["Weak labels", "Some labels are inferred", "Analyst review"],
            ["Random split", "May be optimistic", "Chronological split"],
            ["Partial windows", "Scans are time patterns", "Real rolling aggregation"],
        ],
        0.7,
        1.35,
        12.0,
        3.0,
    )

    slide = blank_slide(prs, "Next Steps")
    add_bullets(
        slide,
        [
            "Collect real SSH, web, and firewall logs from the target environment.",
            "Create analyst-reviewed labels and track false positives.",
            "Implement rolling-window features for SSH brute force, web scanners, firewall scans, and outbound traffic.",
            "Use chronological train/validation/test splits.",
            "Connect ELK export or an API to the hybrid ML scorer.",
        ],
        0.9,
        1.35,
        11.8,
        4.4,
        size=18,
    )

    slide = blank_slide(prs, "Final Takeaway")
    box = slide.shapes.add_textbox(Inches(1.05), Inches(2.0), Inches(11.25), Inches(2.3))
    p = box.text_frame.paragraphs[0]
    p.text = (
        "This project demonstrates a complete cybersecurity log ML pipeline: public and synthetic data, "
        "normalization, EDA, supervised modeling, unified vs specialist comparison, scenario testing, "
        "and ELK ingestion. The next step is validation on real analyst-labeled operational logs."
    )
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = COLORS["ink"]
    p.alignment = PP_ALIGN.CENTER

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT_PATH)
    print(f"wrote {OUTPUT_PATH}")


if __name__ == "__main__":
    build_deck()
